#!/usr/bin/env python3
"""把真实巡检标准 xlsx 转成前端可直接用的 .js 数据文件。

用法（必须用 uv，系统 python3 没有 openpyxl）：
    cd /Users/admin/Code/xunjian-ai-demo
    uv run --with openpyxl python poc/inspection-3d-sandbox/tools/build-items.py

裁剪口径全部在 area-mapping.py 里（唯一需要人工判断的输入）。本脚本只做：
读 xlsx -> 按 area-mapping.py 的口径挑选/折叠/覆盖字段 -> 生成 3 个 ES5 风格
的 .js 文件 -> 跑一遍自检 -> 用 node --check 校验语法。

硬约束：任何和预期结构不符的情况（表头不对、seq 缺失、字段为空、inputType
非法……）都直接抛异常并打印实际读到的内容，不做静默兜底/默认值吞错。
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
from pathlib import Path

import openpyxl

TOOLS_DIR = Path(__file__).resolve().parent


def _load_area_mapping():
    """area-mapping.py 文件名带短横线，不能用普通 import，用 importlib 按路径加载。"""
    spec = importlib.util.spec_from_file_location("area_mapping", TOOLS_DIR / "area-mapping.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


mapping = _load_area_mapping()
PROJECT_ROOT = TOOLS_DIR.parents[2]  # .../xunjian-ai-demo
POC_ROOT = TOOLS_DIR.parent  # .../poc/inspection-3d-sandbox
POC_PARENT = POC_ROOT.parent  # .../poc

# 输出到**两个** POC。
#
# 「沙盘」与「俯视/卫星」是两个刻意互不耦合的独立 POC：各自持有完整的一份数据、契约、
# 组件，运行时零共享，一边改坏了不会波及另一边。但巡检项数据本身必须只有一个真源
# （附件1-3 那份 xlsx + area-mapping.py 的裁剪口径），否则两边会各自漂移出不同的 256 项。
#
# 解法是「构建时同源、运行时零耦合」：本脚本一次跑完写两份独立副本。改 xlsx 或改裁剪
# 口径时跑一条命令同步两边，而不是让两个页面在运行时共享同一个 .js 文件。
DATA_OUT_DIRS = [
    POC_ROOT / "scripts" / "data",
    POC_PARENT / "inspection-3d-aerial" / "scripts" / "data",
]

XLSX_PATH = PROJECT_ROOT / mapping.SOURCE_XLSX_RELATIVE

OUTPUT_GROUPS = {
    "items-entry.js": ["gate"],
    "items-process.js": ["filter", "metering", "regulate", "vent", "blowdown", "launcher"],
    "items-room.js": ["cabinet", "power", "control", "genset", "ups"],
}

VALID_INPUT_TYPES = ("bool", "number")
BOOL_RESULT_TEXT = "合格□ 不合格□"
NUMBER_RESULT_TEXT = "记录数值"


# ==========================================================================
# 1. 读 xlsx
# ==========================================================================

def load_source_rows() -> dict[int, dict]:
    """读「天然气站场」sheet，按 stdSeq(序号) 建索引。

    只对 巡检区域/巡检点项/专业 三列做向下补齐（对应 xlsx 里的合并单元格），
    巡检内容/正确状态/检查结果 不补齐——它们是否为空由 area-mapping.py 里的
    折叠/剔除规则显式处理，不在这里猜测。
    """
    if not XLSX_PATH.is_file():
        raise FileNotFoundError(f"source xlsx not found: {XLSX_PATH}")

    workbook = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if mapping.SOURCE_SHEET not in workbook.sheetnames:
        raise KeyError(
            f"sheet {mapping.SOURCE_SHEET!r} not found, actual sheets: {workbook.sheetnames}"
        )
    sheet = workbook[mapping.SOURCE_SHEET]
    all_rows = list(sheet.iter_rows(values_only=True))

    header_index = None
    for index, row in enumerate(all_rows):
        if row and row[0] == "序号":
            header_index = index
            break
    if header_index is None:
        raise ValueError(f"header row (序号...) not found in sheet {mapping.SOURCE_SHEET!r}")

    header = tuple(all_rows[header_index][: len(mapping.EXPECTED_HEADER)])
    if header != mapping.EXPECTED_HEADER:
        raise ValueError(
            f"unexpected header, expected={mapping.EXPECTED_HEADER!r}, actual={header!r}"
        )

    rows_by_seq: dict[int, dict] = {}
    area = point = discipline = ""
    for row in all_rows[header_index + 1 :]:
        if not row or row[0] is None:
            continue
        raw_seq = row[0]
        if not isinstance(raw_seq, int):
            # 表尾的说明文字行（"说明：具体分区及巡检点项各站结合实际情况调整..."）
            # 序号列不是整数，属于非数据行，直接跳过，不纳入 rows_by_seq。
            continue
        if raw_seq in rows_by_seq:
            raise ValueError(f"duplicate stdSeq in source sheet: {raw_seq}")

        raw_area, raw_point, raw_discipline, raw_content, raw_std, raw_result = row[1:7]
        if raw_area:
            area = str(raw_area).strip()
        if raw_point:
            point = str(raw_point).strip()
        if raw_discipline:
            discipline = str(raw_discipline).strip()

        rows_by_seq[raw_seq] = {
            "area": area,
            "point": point,
            "discipline": discipline,
            "content": (str(raw_content).strip() if raw_content is not None else None),
            "std": (str(raw_std).strip() if raw_std is not None else None),
            "result": (str(raw_result).strip() if raw_result is not None else None),
        }

    return rows_by_seq


# ==========================================================================
# 2. 按 area-mapping.py 的口径构造每个区域的最终条目
# ==========================================================================

def fold_standard(base_std: str, fold_seqs: list[int], rows_by_seq: dict[int, dict]) -> str:
    parts = [base_std]
    for fold_seq in fold_seqs:
        fold_row = rows_by_seq.get(fold_seq)
        if fold_row is None:
            raise KeyError(f"foldSeqs references missing stdSeq: {fold_seq}")
        if fold_row["content"] is not None:
            raise ValueError(
                f"foldSeqs stdSeq={fold_seq} has non-empty content "
                f"({fold_row['content']!r}); folding is only for empty-content "
                "continuation rows, refuse to silently merge a real item"
            )
        if not fold_row["std"]:
            raise ValueError(f"foldSeqs stdSeq={fold_seq} has empty standard, nothing to fold")
        parts.append(fold_row["std"])
    joined = "；".join(p.rstrip("；。") for p in parts)
    return joined


def base_input_fields(result_text: str | None, stdSeq: int) -> dict:
    if result_text == BOOL_RESULT_TEXT:
        return {"inputType": "bool", "value": True, "status": "ok", "unit": None, "thresholdSource": None}
    if result_text == NUMBER_RESULT_TEXT:
        # 真实测量值在 xlsx 里只有占位符（"XX"），没有覆盖表就无法确定 value/unit，
        # 必须显式失败，不允许静默给默认值。
        if stdSeq not in mapping.OVERRIDES:
            raise ValueError(
                f"stdSeq={stdSeq} result column is {NUMBER_RESULT_TEXT!r} but has no "
                "OVERRIDES entry to supply inputType/unit/value"
            )
        return {"inputType": "number", "value": None, "status": "ok", "unit": None, "thresholdSource": None}
    raise ValueError(f"stdSeq={stdSeq} has unexpected result column value: {result_text!r}")


def build_item(
    area_key: str,
    position: int,
    stdSeq: int,
    rows_by_seq: dict[int, dict],
    allowed_source_areas: list[str],
    fold_seqs: list[int] | None = None,
    merged_from: list[str] | None = None,
    borrowed_from: str | None = None,
) -> dict:
    row = rows_by_seq.get(stdSeq)
    if row is None:
        raise KeyError(f"{area_key}: stdSeq={stdSeq} not found in source rows")
    if row["area"] not in allowed_source_areas:
        raise ValueError(
            f"{area_key}: stdSeq={stdSeq} belongs to xlsx area {row['area']!r}, "
            f"expected one of {allowed_source_areas!r}"
        )

    title = row["content"]
    standard = row["std"]
    if fold_seqs:
        if title is None:
            raise ValueError(f"{area_key}: stdSeq={stdSeq} has empty content, cannot be a fold anchor")
        standard = fold_standard(standard, fold_seqs, rows_by_seq)

    fields = base_input_fields(row["result"], stdSeq)

    item = {
        "id": f"{area_key}-{position}",
        "seq": position,
        "areaKey": area_key,
        "point": row["point"],
        "discipline": row["discipline"],
        "title": title,
        "standard": standard,
        "inputType": fields["inputType"],
        "unit": fields["unit"],
        "value": fields["value"],
        "status": fields["status"],
        "tag": None,
        "stdSeq": stdSeq,
        "mergedFrom": merged_from,
        "borrowedFrom": borrowed_from,
        "thresholdSource": fields["thresholdSource"],
        # 数值型量程：见下方 NUMBER_RANGES 的强制校验。bool 项恒为 None。
        "min": None,
        "max": None,
        "rangeSource": None,
    }

    override = mapping.OVERRIDES.get(stdSeq)
    if override:
        for key, value in override.items():
            if key not in ("content", "tag", "value", "status", "inputType", "unit", "thresholdSource"):
                raise KeyError(f"{area_key}: stdSeq={stdSeq} override has unknown key {key!r}")
            if key == "content":
                item["title"] = value
            else:
                item[key] = value

    if not item["title"]:
        raise ValueError(f"{area_key}: stdSeq={stdSeq} resolved to empty title")
    if not item["standard"]:
        raise ValueError(f"{area_key}: stdSeq={stdSeq} resolved to empty standard")
    if item["inputType"] not in VALID_INPUT_TYPES:
        raise ValueError(f"{area_key}: stdSeq={stdSeq} invalid inputType {item['inputType']!r}")
    if item["inputType"] == "number":
        if item["value"] is None:
            raise ValueError(f"{area_key}: stdSeq={stdSeq} number item has no value")
        if item["unit"] not in mapping.VALID_UNITS:
            raise ValueError(f"{area_key}: stdSeq={stdSeq} number item has invalid unit {item['unit']!r}")

        # 量程必须由 NUMBER_RANGES 显式给出，缺了直接抛错、不给默认值。
        #
        # 为什么这条要硬：巡检项列表把数值型渲染成 role="spinbutton"，ARIA 强制要求
        # aria-valuemin/aria-valuemax。数据层不给区间，UI 层就只能自己编一个"通用
        # 量程"——页面上会出现一个看起来权威、实际是编的数字，而且没有任何机制会
        # 提示它是编的。所以把这份责任钉在数据层，并在构建期强制。
        rng = mapping.NUMBER_RANGES.get(stdSeq)
        if rng is None:
            raise ValueError(
                f"{area_key}: stdSeq={stdSeq} 是 number 型但 NUMBER_RANGES 里没有量程条目，"
                f"请在 area-mapping.py 的 NUMBER_RANGES 里补上 min/max/rangeSource"
                f"（title={item['title']!r} unit={item['unit']!r}）"
            )
        for key in ("min", "max", "rangeSource"):
            if key not in rng:
                raise KeyError(f"{area_key}: stdSeq={stdSeq} NUMBER_RANGES 条目缺少 {key!r}")
        if not isinstance(rng["min"], (int, float)) or not isinstance(rng["max"], (int, float)):
            raise TypeError(f"{area_key}: stdSeq={stdSeq} NUMBER_RANGES 的 min/max 必须是数字")
        if rng["min"] >= rng["max"]:
            raise ValueError(
                f"{area_key}: stdSeq={stdSeq} NUMBER_RANGES 的 min({rng['min']}) 必须小于 max({rng['max']})"
            )
        item["min"] = rng["min"]
        item["max"] = rng["max"]
        item["rangeSource"] = rng["rangeSource"]

        # 注意：**不校验 value 落在 [min,max] 内**。剧本 B/C 的两条异常项恰恰是
        # 越界读数（gate-40 的 3 MPa 低于下限 4.0、filter-4 的 0.14 MPa 高于上限
        # 0.1），越界是它们成为异常项的原因。这里只保证区间自身自洽。
    else:
        if item["unit"] is not None:
            raise ValueError(f"{area_key}: stdSeq={stdSeq} bool item must not have a unit")
        if mapping.NUMBER_RANGES.get(stdSeq) is not None:
            raise ValueError(
                f"{area_key}: stdSeq={stdSeq} 是 bool 型，却在 NUMBER_RANGES 里有量程条目——"
                f"很可能是 OVERRIDES 里漏了把它改成 number 型"
            )

    return item


def build_area_items(area_key: str, area_def: dict, rows_by_seq: dict[int, dict]) -> list[dict]:
    allowed_source_areas = area_def["sourceAreas"]
    specs: list[dict] = []

    if area_key == "cabinet":
        fold_map: dict[int, list[int]] = area_def["fold"]
        drop_seqs: set[int] = area_def["dropSeqs"]
        for stdSeq in area_def["sourceRange"]:
            if stdSeq in drop_seqs:
                continue
            spec = {"seq": stdSeq}
            if stdSeq in fold_map:
                spec["foldSeqs"] = fold_map[stdSeq]
            specs.append(spec)
    else:
        specs.extend(area_def["sequence"])

    items: list[dict] = []
    for position, spec in enumerate(specs, start=1):
        items.append(
            build_item(
                area_key,
                position,
                spec["seq"],
                rows_by_seq,
                allowed_source_areas,
                fold_seqs=spec.get("foldSeqs"),
                merged_from=spec.get("mergedFrom"),
            )
        )

    borrow = area_def.get("borrow")
    if borrow:
        position = len(items) + 1
        items.append(
            build_item(
                area_key,
                position,
                borrow["seq"],
                rows_by_seq,
                [borrow["sourceArea"]],
                borrowed_from=borrow["sourceArea"],
            )
        )

    return items


# ==========================================================================
# 3. 自检
# ==========================================================================

def self_check(all_items: dict[str, list[dict]]) -> None:
    grand_total = 0
    for area_key in mapping.AREA_ORDER:
        items = all_items[area_key]
        target = mapping.AREA_DEFS[area_key]["targetCount"]
        if len(items) != target:
            raise AssertionError(f"{area_key}: expected {target} items, got {len(items)}")
        grand_total += len(items)

        seqs_seen = [item["seq"] for item in items]
        if seqs_seen != list(range(1, len(items) + 1)):
            raise AssertionError(f"{area_key}: seq must be 1..{len(items)} contiguous, got {seqs_seen}")

        ids_seen = {item["id"] for item in items}
        if len(ids_seen) != len(items):
            raise AssertionError(f"{area_key}: duplicate id detected")

        for item in items:
            if not item["title"] or not isinstance(item["title"], str):
                raise AssertionError(f"{area_key}: item {item['id']} has empty/invalid title")
            if not item["standard"] or not isinstance(item["standard"], str):
                raise AssertionError(f"{area_key}: item {item['id']} has empty/invalid standard")
            if item["inputType"] not in VALID_INPUT_TYPES:
                raise AssertionError(f"{area_key}: item {item['id']} invalid inputType")
            if item["inputType"] == "number" and item["unit"] not in mapping.VALID_UNITS:
                raise AssertionError(f"{area_key}: item {item['id']} number item invalid unit")

    if grand_total != 256:
        raise AssertionError(f"grand total must be 256, got {grand_total}")

    print(f"[self-check] OK, 12 areas, grand total = {grand_total}")


# ==========================================================================
# 4. 生成 .js
# ==========================================================================

def render_js(
    source_areas: list[str],
    items_by_area: dict[str, list[dict]],
    seed_order: tuple[str, ...] | None = None,
) -> str:
    lines = []
    lines.append("// 本文件由 tools/build-items.py 生成，不要手改；改数据请改 xlsx 或 area-mapping.py 后重跑。")
    lines.append(f"// 源文件：{Path(mapping.SOURCE_XLSX_RELATIVE).name}（sheet: {mapping.SOURCE_SHEET}）")
    counts = ", ".join(f"{area_key}={len(items_by_area[area_key])}" for area_key in source_areas)
    lines.append(f"// 条目数：{counts}")
    lines.append("(function () {")
    lines.append('  "use strict";')
    lines.append("")

    if seed_order is not None:
        # 键顺序纪律：window.DemoItems 的键顺序必须等于 Map3DContract.AREA_IDS（也就是真实
        # App「选择区域」列表的顺序），因为 Map3DContract.assertIdSet 是**含顺序**全等校验。
        #
        # 而三个 items-*.js 是按「工艺区 / 房间区」分组切分的，加载顺序与区域语义顺序天然
        # 不一致（launcher 收发球区属工艺区、排在 items-process.js 里，却应位于 cabinet 之后）。
        # 若各文件各自 `window.DemoItems.<key> = ...`，键顺序由加载顺序决定，assertIdSet 会红。
        # 所以在最先加载的本文件里一次性按契约顺序预置 12 个键，后续文件只做赋值、不新增键。
        lines.append("  // 键顺序 = Map3DContract.AREA_IDS（含顺序全等校验，见 map3d/contract.js 的 assertIdSet）")
        lines.append("  if (window.DemoItems) {")
        lines.append('    throw new Error("items-entry.js 必须是三个 items-*.js 中最先加载的一个：'
                     'window.DemoItems 已存在，说明 index.html 的 script 顺序被改动过");')
        lines.append("  }")
        lines.append("  window.DemoItems = {")
        seeded = ", ".join(f"{area_key}: null" for area_key in seed_order)
        lines.append(f"    {seeded}")
        lines.append("  };")
        lines.append("")
    else:
        lines.append("  if (!window.DemoItems) {")
        lines.append('    throw new Error("本文件必须在 items-entry.js 之后加载：window.DemoItems 尚未预置键顺序");')
        lines.append("  }")
        lines.append("")
    for area_key in source_areas:
        payload = json.dumps(items_by_area[area_key], ensure_ascii=False, indent=2)
        # json.dumps 输出的是合法的 JS 数组/对象字面量（true/false/null 与 JS 一致），
        # 缩进整体往右挪 2 格以贴合 IIFE 内部的书写风格。
        indented_payload = "\n".join(
            ("  " + line if line else line) for line in payload.splitlines()
        )
        lines.append(f"  window.DemoItems.{area_key} = {indented_payload};")
        lines.append("")
    lines.append("})();")
    lines.append("")
    return "\n".join(lines)


def write_output_file(
    filename: str,
    area_keys: list[str],
    items_by_area: dict[str, list[dict]],
    seed_order: tuple[str, ...] | None = None,
) -> list[Path]:
    """写入全部 DATA_OUT_DIRS，返回所有落盘路径。"""
    payload = render_js(area_keys, items_by_area, seed_order)
    written = []
    for out_dir in DATA_OUT_DIRS:
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / filename
        out_path.write_text(payload, encoding="utf-8")
        written.append(out_path)
    return written


def run_node_check(path: Path) -> None:
    result = subprocess.run(["node", "--check", str(path)], capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"node --check failed for {path}:\n{result.stdout}\n{result.stderr}")
    print(f"[node --check] OK  {path.relative_to(PROJECT_ROOT)}")


# ==========================================================================
# 5. 日志：裁剪/去重/脏数据明细
# ==========================================================================

def print_cut_log() -> None:
    print("\n" + "=" * 78)
    print("裁剪 / 去重 / 脏数据处理明细")
    print("=" * 78)
    for area_key in mapping.AREA_ORDER:
        area_def = mapping.AREA_DEFS[area_key]
        dropped = area_def["dropped"]
        print(f"\n[{area_key}] {area_def['cn']}  目标 {area_def['targetCount']} 项")
        print(f"  裁剪口径：{area_def['cutRule']}")
        if not dropped:
            print("  未裁剪任何条目。")
            continue
        print(f"  共 {len(dropped)} 条未进入最终列表：")
        for seq, src_area, point, content, reason, note in dropped:
            print(f"    - stdSeq={seq:<4} [{src_area}] 巡检点项={point!r} 巡检内容={content!r}")
            print(f"        reason={reason}  {note}")


# ==========================================================================
# main
# ==========================================================================

def main() -> None:
    rows_by_seq = load_source_rows()
    print(f"[load] {XLSX_PATH.name} / sheet={mapping.SOURCE_SHEET} -> {len(rows_by_seq)} stdSeq rows")

    items_by_area: dict[str, list[dict]] = {}
    for area_key in mapping.AREA_ORDER:
        items_by_area[area_key] = build_area_items(area_key, mapping.AREA_DEFS[area_key], rows_by_seq)
        area_def = mapping.AREA_DEFS[area_key]
        print(
            f"[build] {area_key:<10} {area_def['cn']:<8} "
            f"kept={len(items_by_area[area_key])} target={area_def['targetCount']} "
            f"dropped={len(area_def['dropped'])}"
        )

    self_check(items_by_area)

    written_paths = []
    # OUTPUT_GROUPS 的第一项（items-entry.js）负责按 AREA_ORDER 预置全部 12 个键，
    # 详见 render_js 里「键顺序纪律」那段注释。
    first_filename = next(iter(OUTPUT_GROUPS))
    for filename, area_keys in OUTPUT_GROUPS.items():
        seed_order = mapping.AREA_ORDER if filename == first_filename else None
        out_paths = write_output_file(filename, area_keys, items_by_area, seed_order)
        written_paths.extend(out_paths)
        area_summary = ", ".join(f"{k}={len(items_by_area[k])}" for k in area_keys)
        for out_path in out_paths:
            print(f"[write] {out_path.relative_to(PROJECT_ROOT)}  ({area_summary})")

    for out_path in written_paths:
        run_node_check(out_path)

    print_cut_log()

    print("\n" + "=" * 78)
    print(f"DONE. 12 areas / 256 items written to {len(DATA_OUT_DIRS)} POC(s):")
    for out_dir in DATA_OUT_DIRS:
        print("  -", out_dir.relative_to(PROJECT_ROOT))
    print("=" * 78)


if __name__ == "__main__":
    main()
